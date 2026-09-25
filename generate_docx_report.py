import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_report():
    doc = Document()
    
    # Page setup - Margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    # Styling helper functions
    def set_cell_background(cell, fill_hex):
        shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        cell._tc.get_or_add_tcPr().append(shading_elm)
        
    def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
        tcPr.append(tcMar)

    def add_title(text, size=24, color=RGBColor(14, 116, 144)):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.bold = True
        run.font.name = 'Arial'
        run.font.size = Pt(size)
        run.font.color.rgb = color
        p.paragraph_format.space_after = Pt(12)
        return p

    def add_heading1(text):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = True
        run.font.name = 'Arial'
        run.font.size = Pt(15)
        run.font.color.rgb = RGBColor(14, 116, 144) # Deep Teal / Blue
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        return p

    def add_heading2(text):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = True
        run.font.name = 'Arial'
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(30, 41, 59)
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        return p

    def add_paragraph(text, space_after=6):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(51, 65, 85)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(space_after)
        return p

    def add_code_block(code_text):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        set_cell_background(cell, "F8FAFC")
        set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
        
        # Border
        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/><w:left w:val="single" w:sz="16" w:space="0" w:color="0284C7"/><w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/><w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/></w:tcBorders>')
        tcPr.append(borders)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(code_text)
        run.font.name = 'Consolas'
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(30, 41, 59)
        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # ==========================
    # PAGE 1: COVER PAGE
    # ==========================
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_inst = p_inst.add_run("SRI SHAKTHI INSTITUTE OF ENGINEERING AND TECHNOLOGY\n")
    r_inst.bold = True
    r_inst.font.name = 'Arial'
    r_inst.font.size = Pt(14)
    r_inst.font.color.rgb = RGBColor(15, 23, 42)
    
    r_sub = p_inst.add_run("(An Autonomous Institution | Approved by AICTE, Affiliated to Anna University)\nCoimbatore – 641062, Tamil Nadu\n\n\n\n\n")
    r_sub.font.name = 'Calibri'
    r_sub.font.size = Pt(11)
    r_sub.font.color.rgb = RGBColor(100, 116, 139)

    p_proj = doc.add_paragraph()
    p_proj.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_proj_tag = p_proj.add_run("MERN PROJECT\n\n")
    r_proj_tag.bold = True
    r_proj_tag.font.name = 'Arial'
    r_proj_tag.font.size = Pt(18)
    r_proj_tag.font.color.rgb = RGBColor(14, 116, 144)
    
    r_proj_name = p_proj.add_run("ResQ: INTELLIGENT DISASTER RESPONSE &\nRESCUE PLATFORM\n\n\n\n\n\n\n")
    r_proj_name.bold = True
    r_proj_name.font.name = 'Arial'
    r_proj_name.font.size = Pt(22)
    r_proj_name.font.color.rgb = RGBColor(15, 23, 42)

    p_done = doc.add_paragraph()
    p_done.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_done_lbl = p_done.add_run("DONE BY:\n")
    r_done_lbl.bold = True
    r_done_lbl.font.name = 'Arial'
    r_done_lbl.font.size = Pt(11)
    r_done_lbl.font.color.rgb = RGBColor(71, 85, 105)
    
    r_name = p_done.add_run("NAVANEETHAN S\n")
    r_name.bold = True
    r_name.font.name = 'Arial'
    r_name.font.size = Pt(13)
    r_name.font.color.rgb = RGBColor(15, 23, 42)

    doc.add_page_break()

    # ==========================
    # PAGE 2: ABSTRACT & TECH STACK
    # ==========================
    add_heading1("ABSTRACT")
    add_paragraph("The ResQ platform is a comprehensive full-stack disaster management and emergency coordination ecosystem engineered using MongoDB, Express.js, React.js, and Node.js (MERN stack). Designed to eliminate communication delays and operational chaos during life-critical natural and human-made crises (floods, earthquakes, wildfires, industrial accidents, and landslides), ResQ bridges the divide between trapped citizens, volunteer rescue squads, emergency first responders, and central disaster management authorities.")
    add_paragraph("The platform is built around an automated, browser-level GPS location acquisition engine paired with OpenStreetMap reverse-geocoding, ensuring that emergency reports and SOS broadcasts are tagged with pinpoint latitude and longitude telemetry. Built on a low-latency WebSocket communication backbone powered by Socket.IO, critical SOS beacons and disaster alerts are broadcast instantaneously across all active rescue squad terminals without requiring manual page refreshes.")
    add_paragraph("In addition to its real-time incident reporting features, ResQ incorporates an interactive multi-layer GIS mapping interface offering Google Satellite/Hybrid, Google Streets, and CartoDB Dark Matter views with integrated one-tap Google Maps turn-by-turn navigation routing. Dedicated portals are customized for four distinct operational roles: Citizens, Volunteers, Rescue Squads, and System Administrators. By combining modern React components, high-throughput Node.js microservices, and geospatial 2dsphere indexing in MongoDB, ResQ provides a scalable, high-resilience digital foundation for emergency rescue operations.")

    add_heading1("TECH STACK")
    
    add_heading2("Frontend")
    add_paragraph("The client-side application of ResQ is developed using React.js 18 with Vite, delivering an ultra-responsive, component-driven Single-Page Application (SPA). React governs all operational views including the public landing page, role-specific authentication portals, the Citizen safety command center, the Volunteer mission dashboard, the Rescue squad dispatch terminal, and the Admin oversight center.")
    add_paragraph("Tailwind CSS powers a dark glassmorphism visual design system inspired by mission-critical control centers (#0a0d14 background, neon blue #00d4ff and emergency red #ff3b3b accents, frosted glass panels, and smooth Framer Motion transitions). The interactive GIS interface is driven by React-Leaflet, rendering customizable tile sets from Google Maps (Satellite Hybrid, Streets, Terrain) and CartoDB Dark Matter with dynamic marker popups. Axios is utilized for secure, intercepted REST API communication with automated JWT bearer token management.")

    add_heading2("Backend")
    add_paragraph("The backend architecture is built with Node.js and Express.js, providing an asynchronous, non-blocking REST API gateway and WebSocket server. The backend follows a strict modular structure separating routes, controllers, middleware, and Mongoose models.")
    add_paragraph("Socket.IO orchestrates real-time bidirectional event distribution across multiple pub/sub channels (sos-alert, new-alert, rescue-update, and volunteer-location). Security is enforced using JSON Web Tokens (JWT) for stateless authentication, bcryptjs for cryptographic password hashing, Helmet for HTTP header protection, and custom role-based access control (RBAC) middleware verifying user authorizations.")

    add_heading2("Database")
    add_paragraph("ResQ utilizes MongoDB as its primary document-oriented NoSQL database. MongoDB's schema flexibility easily accommodates polymorphic incident structures, real-time telemetry coordinates, volunteer capability rosters, and crisis inventory logs. Mongoose serves as the Object Data Modeling (ODM) layer, enforcing validation and schema integrity. Spherical 2dsphere spatial indexes are applied to geographical coordinate pairs [longitude, latitude] across DisasterReport, SOSRequest, Resource, and Alert collections, enabling rapid geospatial proximity queries.")

    doc.add_page_break()

    # ==========================
    # PAGE 3: MODULES
    # ==========================
    add_heading1("MODULES")

    add_heading2("Module 1: User Authentication & Role-Based Access Control (RBAC)")
    add_paragraph("Manages user registration, credential verification, and account permissions for four distinct operational tiers: Citizen, Volunteer, Rescue Team Member, and Administrator. Utilizes bcryptjs (10 salt rounds) for password hashing and issues signed JWT bearer tokens for stateless session verification. Role-based middleware guards ensure sensitive operations (such as squad dispatch and resource allocation) remain restricted.")

    add_heading2("Module 2: Disaster & Hazard Incident Reporting")
    add_paragraph("Enables citizens and field scouts to report disasters categorized under flood, earthquake, fire, cyclone, accident, landslide, tsunami, or other hazards. Features automated browser GPS telemetry locking with reverse geocoding via OpenStreetMap Nominatim, allowing precise address resolution without requiring manual user input.")

    add_heading2("Module 3: Instant One-Tap Emergency SOS Broadcast")
    add_paragraph("Engineered for life-or-death situations, this module provides an immediate one-tap SOS beacon that captures the victim's exact GPS coordinates and broadcasts a high-priority WebSocket alert (sos-alert) to all connected rescue dispatch units and nearby field responders within milliseconds.")

    add_heading2("Module 4: Tactical Multi-Layer Live GIS Map")
    add_paragraph("A full-screen interactive command map rendering active disasters, emergency SOS markers, volunteer positions, hospitals, and relief shelters. Features user-selectable basemaps including Google Satellite Hybrid, Google Streets, and CartoDB Dark Matter. Every incident pin provides one-click deep links to open the exact coordinates in Google Maps or launch Google Maps turn-by-turn navigation.")

    add_heading2("Module 5: Volunteer Squads & Mission Mobilization")
    add_paragraph("Allows registered community volunteers to view active distress calls, enroll in relief missions, form specialized response squads, track rescue progress, and accumulate mission completion badges and volunteer rating points.")

    add_heading2("Module 6: Rescue Squad Emergency Command & Field Navigation")
    add_paragraph("Provides dedicated rescue personnel with prioritized triage feeds (Critical, High, Medium, Low). Rescue squads can accept emergency calls, update mission status (pending -> accepted -> on_the_way -> rescued -> closed), and launch turn-by-turn navigation directly to victims in need.")

    add_heading2("Module 7: Emergency Relief Resource & Shelter Management")
    add_paragraph("Maintains real-time inventory tracking for critical survival supplies across categories: medical kits, trauma blood bags, drinking water, emergency food rations, lifeboats, and shelter beds. Features status monitors (available, low, depleted) to prevent resource stockouts during operations.")

    add_heading2("Module 8: AI-Assisted Severity Scoring & Verification")
    add_paragraph("Integrates an AI analysis engine that parses disaster descriptions, affected casualty numbers, and incident categories to generate automated severity scores (0-100), risk tiers, predicted impact populations, required supply checklists, and false-alarm detection scores.")

    add_heading2("Module 9: Central Admin Command Center & Telemetry Analytics")
    add_paragraph("Provides comprehensive platform oversight for disaster management authorities. Visualizes system-wide rescue statistics, total resolved incidents, volunteer distribution density, and audit logs.")

    add_heading2("Module Integration & Operational Workflow")
    add_paragraph("The modules operate as a synchronized emergency coordination network. The operational workflow proceeds as follows:")
    add_paragraph("Citizen Report / Emergency SOS -> GPS Auto-Lock & Telemetry -> Real-Time Socket.IO Alert Broadcast -> Live GIS Map Update -> Rescue Squad Triage & Dispatch -> Turn-by-Turn Navigation via Google Maps -> Volunteer Support & Resource Mobilization -> Incident Status Resolution -> Central Admin Audit.")

    doc.add_page_break()

    # ==========================
    # PAGE 4: FRONTEND CODE
    # ==========================
    add_heading1("FRONTEND CODE")

    add_heading2("index.html:")
    add_code_block("""<!doctype html>
<html lang="en">
  <head>
    <meta charset="UTF-8" />
    <link rel="icon" type="image/svg+xml" href="/vite.svg" />
    <meta name="viewport" content="width=device-width, initial-scale=1.0" />
    <title>ResQ - Intelligent Disaster Response & Rescue Platform</title>
    <meta name="description" content="AI-Powered Emergency Platform for Disaster Response" />
    <link href="https://fonts.googleapis.com/css2?family=Inter:wght@300;400;500;600;700&display=swap" rel="stylesheet" />
    <link rel="stylesheet" href="https://unpkg.com/leaflet@1.9.4/dist/leaflet.css" />
  </head>
  <body class="bg-[#0a0d14] text-white font-sans overflow-x-hidden">
    <div id="root"></div>
    <script type="module" src="/src/main.jsx"></script>
  </body>
</html>""")

    add_heading2("tailwind.config.js:")
    add_code_block("""/** @type {import('tailwindcss').Config} */
export default {
  content: ["./index.html", "./src/**/*.{js,ts,jsx,tsx}"],
  theme: {
    extend: {
      colors: {
        primary: '#00d4ff',
        secondary: '#0a0f1c',
        accent: '#ff3b3b',
        dark: {
          100: '#1a1f2e',
          200: '#141825',
          300: '#0f1219',
          400: '#0a0d14'
        },
        neon: {
          blue: '#00d4ff',
          purple: '#a855f7',
          green: '#22c55e'
        }
      },
      fontFamily: {
        sans: ['Inter', 'sans-serif']
      }
    }
  },
  plugins: []
}""")

    add_heading2("LiveMap.jsx (Sample Component Excerpt):")
    add_code_block("""import React, { useState, useEffect } from 'react';
import { MapContainer, TileLayer, Marker, Popup, useMap } from 'react-leaflet';
import L from 'leaflet';
import { FiActivity, FiUsers, FiCrosshair, FiSearch, FiExternalLink, FiNavigation } from 'react-icons/fi';
import 'leaflet/dist/leaflet.css';
import { detectCurrentLocation } from '../utils/location';

const TILE_LAYERS = {
  googleHybrid: {
    id: 'googleHybrid',
    name: 'Google Satellite',
    url: 'https://{s}.google.com/vt/lyrs=y&x={x}&y={y}&z={z}',
    subdomains: ['mt0', 'mt1', 'mt2', 'mt3']
  },
  googleStreets: {
    id: 'googleStreets',
    name: 'Google Streets',
    url: 'https://{s}.google.com/vt/lyrs=m&x={x}&y={y}&z={z}',
    subdomains: ['mt0', 'mt1', 'mt2', 'mt3']
  }
};

export default function LiveMap() {
  const [activeLayer, setActiveLayer] = useState('googleHybrid');
  const [userLoc, setUserLoc] = useState(null);

  useEffect(() => {
    detectCurrentLocation().then(loc => setUserLoc([loc.lat, loc.lng]));
  }, []);

  return (
    <div className="w-full h-[calc(100vh-8rem)] min-h-[600px] relative rounded-3xl overflow-hidden border border-white/10">
      <MapContainer center={userLoc || [20.5937, 78.9629]} zoom={13} style={{ height: '100%', width: '100%' }}>
        <TileLayer url={TILE_LAYERS[activeLayer].url} subdomains={TILE_LAYERS[activeLayer].subdomains} />
        {userLoc && (
          <Marker position={userLoc}>
            <Popup>
              <div className="p-2 text-dark-400">
                <h4 className="font-bold">Your Current GPS</h4>
                <a href={`https://www.google.com/maps/dir/?api=1&destination=${userLoc[0]},${userLoc[1]}`} 
                   target="_blank" rel="noreferrer" className="text-primary font-bold">
                  Navigate via Google Maps →
                </a>
              </div>
            </Popup>
          </Marker>
        )}
      </MapContainer>
    </div>
  );
}""")

    doc.add_page_break()

    # ==========================
    # PAGE 5: BACKEND CODE
    # ==========================
    add_heading1("BACKEND CODE")

    add_heading2("server.js (Main Entry Point):")
    add_code_block("""require('dotenv').config();
const express = require('express');
const cors = require('cors');
const helmet = require('helmet');
const http = require('http');
const { Server } = require('socket.io');
const connectDB = require('./config/db');

const app = express();
const server = http.createServer(app);
const io = new Server(server, {
  cors: { origin: process.env.CLIENT_URL || '*', methods: ['GET', 'POST', 'PUT', 'DELETE'] }
});

app.use(express.json());
app.use(cors());
app.use(helmet());

connectDB();
app.set('io', io);

io.on('connection', (socket) => {
  socket.on('sos-alert', (data) => io.emit('sos-alert', data));
  socket.on('rescue-update', (data) => io.emit('rescue-update', data));
});

app.use('/api/auth', require('./routes/authRoutes'));
app.use('/api/disasters', require('./routes/disasterRoutes'));
app.use('/api/sos', require('./routes/sosRoutes'));
app.use('/api/volunteers', require('./routes/volunteerRoutes'));
app.use('/api/resources', require('./routes/resourceRoutes'));

const PORT = process.env.PORT || 5000;
server.listen(PORT, () => console.log(`Server running on port ${PORT}`));""")

    add_heading2("DisasterReport Model:")
    add_code_block("""const mongoose = require('mongoose');

const disasterReportSchema = new mongoose.Schema({
  userId: { type: mongoose.Schema.Types.ObjectId, ref: 'User' },
  type: { 
    type: String, 
    enum: ['flood', 'earthquake', 'fire', 'cyclone', 'accident', 'landslide', 'tsunami', 'other'] 
  },
  description: { type: String },
  images: [{ type: String }],
  location: {
    type: { type: String, default: 'Point' },
    coordinates: [Number], // [longitude, latitude]
    address: String
  },
  severity: { type: String, enum: ['low', 'medium', 'high', 'critical'] },
  aiScore: { type: Number, min: 0, max: 100 },
  status: { type: String, enum: ['reported', 'verified', 'responding', 'resolved'], default: 'reported' },
  affectedCount: { type: Number },
  createdAt: { type: Date, default: Date.now }
});

disasterReportSchema.index({ location: '2dsphere' });
module.exports = mongoose.model('DisasterReport', disasterReportSchema);""")

    add_heading2("SOSRequest Model:")
    add_code_block("""const mongoose = require('mongoose');

const sosRequestSchema = new mongoose.Schema({
  citizen: { type: mongoose.Schema.Types.ObjectId, ref: 'User' },
  location: {
    type: { type: String, default: 'Point' },
    coordinates: [Number],
    address: String
  },
  message: { type: String },
  priority: { type: String, enum: ['low', 'medium', 'high', 'critical'], default: 'high' },
  assignedTeam: { type: mongoose.Schema.Types.ObjectId, ref: 'User' },
  status: { type: String, enum: ['pending', 'accepted', 'on_the_way', 'rescued', 'closed'], default: 'pending' },
  responseTime: { type: Number },
  createdAt: { type: Date, default: Date.now }
});

sosRequestSchema.index({ location: '2dsphere' });
module.exports = mongoose.model('SOSRequest', sosRequestSchema);""")

    add_heading2("SOS Controller & Dispatch Route:")
    add_code_block("""const SOSRequest = require('../models/SOSRequest');

exports.createSOS = async (req, res) => {
  try {
    const sosData = { ...req.body };
    if (req.user) sosData.citizen = req.user._id;

    const sos = await SOSRequest.create(sosData);
    
    // Broadcast via WebSockets to all rescue squads
    const io = req.app.get('io');
    if (io) {
      io.emit('sos-alert', sos);
    }
    
    res.status(201).json({ success: true, data: sos });
  } catch (error) {
    res.status(500).json({ success: false, message: error.message });
  }
};""")

    doc.add_page_break()

    # ==========================
    # PAGE 6: SYSTEM OUTPUT
    # ==========================
    add_heading1("OUTPUT")

    outputs = [
        ("Fig 1: Landing Page & Emergency Gateway", "The public landing view featuring distinct operational pathways: 'Get Started' for registered personnel and an immediate red-pulsing 'EMERGENCY SOS' button providing zero-login crisis broadcasting for citizens."),
        ("Fig 2: Role-Based Authentication Portal", "Secure login and registration interface allowing users to authenticate under their specific role: Citizen, Volunteer, Rescue Team Member, or System Administrator, with encrypted JWT generation."),
        ("Fig 3: Citizen Dashboard", "The citizen's operational dashboard displaying regional hazard alerts, real-time safety instructions, an incident status feed, and a quick SOS trigger button."),
        ("Fig 4: Incident Report Submission with Auto-GPS", "The incident reporting form with integrated browser GPS auto-detection, resolving the user's exact geographical coordinates and reverse-geocoded physical address automatically."),
        ("Fig 5: Interactive Multi-Layer GIS Live Map", "Full-screen tactical map rendering live incidents, volunteer locations, hospitals, and shelters with Google Satellite Hybrid and Google Streets tiles, plus one-click Google Maps Navigation buttons."),
        ("Fig 6: Critical Emergency SOS Broadcast", "High-priority emergency beacon display transmitting live victim coordinates over WebSockets to all rescue dispatch consoles with acoustic and visual ping indicators."),
        ("Fig 7: Volunteer Mission Control & Squad Mobilization", "Specialized volunteer interface for browsing active emergencies, joining rescue squads, tracking completed missions, and managing responder badges."),
        ("Fig 8: Rescue Squad Command Center & Field Navigation", "Emergency triage console prioritizing incidents by severity, providing squad assignment controls and one-click turn-by-turn navigation via Google Maps."),
        ("Fig 9: Emergency Resource & Shelter Inventory", "Supply management registry tracking hospital bed availability, blood bank stocks, medical kits, and relief shelter capacity with depletion threshold indicators."),
        ("Fig 10: Admin Central Command & Telemetry Analytics", "Comprehensive analytical overview displaying platform-wide incident resolution rates, responder density, and disaster classification distributions.")
    ]

    for title, desc in outputs:
        p_fig = doc.add_paragraph()
        run_fig = p_fig.add_run(f"\n{title}\n")
        run_fig.bold = True
        run_fig.font.name = 'Arial'
        run_fig.font.size = Pt(11.5)
        run_fig.font.color.rgb = RGBColor(14, 116, 144)
        
        p_desc = doc.add_paragraph()
        run_desc = p_desc.add_run(desc)
        run_desc.font.name = 'Calibri'
        run_desc.font.size = Pt(10.5)
        run_desc.font.color.rgb = RGBColor(51, 65, 85)
        p_desc.paragraph_format.space_after = Pt(8)

    doc.add_page_break()

    # ==========================
    # PAGE 7: CONCLUSION & FUTURE ENHANCEMENTS
    # ==========================
    add_heading1("CONCLUSION & FUTURE ENHANCEMENTS")

    add_heading2("Conclusion")
    add_paragraph("The ResQ – Intelligent Disaster Response & Rescue Platform successfully demonstrates the efficacy of modern MERN stack technologies in solving critical communication and coordination challenges during emergency scenarios. By combining real-time WebSocket broadcasting with automated GPS location acquisition, multi-layer Google Maps GIS visualization, and role-specific workflows (Citizen, Volunteer, Rescue Team, Admin), ResQ dramatically reduces disaster triage response times and empowers coordinated life-saving relief efforts.")

    add_heading2("Future Enhancements")
    add_paragraph("1. Offline-First PWA Mesh Networking: Implementing Progressive Web App service workers and Bluetooth Low Energy (BLE) / Wi-Fi Direct peer-to-peer mesh protocols to enable emergency communication in environments where cellular telecommunications towers are down.")
    add_paragraph("2. UAV & Drone Video Telemetry Ingestion: Streaming live aerial drone feeds directly into Leaflet map overlays for real-time flood boundary tracking, wildfire progression monitoring, and structural damage assessment.")
    add_paragraph("3. Automated Cellular SMS / USSD Fallback Gateway: Integrating cellular gateways (e.g. Twilio) allowing citizens without smartphones or internet access to trigger emergency SOS beacons and report hazards via standard text messages.")
    add_paragraph("4. Predictive AI Meteorological Modeling: Coupling machine learning models with regional weather and river gauge APIs to predict flash floods and provide predictive evacuation warnings hours before disasters strike.")

    # Save document
    output_path = r"c:\project\Mern.web\ResQ_Project_Report.docx"
    doc.save(output_path)
    print(f"Report successfully saved to: {output_path}")

if __name__ == "__main__":
    create_report()
